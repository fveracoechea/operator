"""Read and write the Performance Highlights workbook.

  workbook.py export <workbook.xlsx> <data.json>
  workbook.py build  <data.json> <workbook.xlsx>
"""
import json, re, sys, datetime as dt
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
from openpyxl.workbook.properties import CalcProperties

FIELDS = ["rank", "title", "category", "repos", "start", "latest", "status", "decision", "result",
          "why", "cross_team", "evidence", "main_link", "cross_team_link"]
HEADERS = ["#", "Highlight", "Category", "Repos and teams", "Start", "Latest", "Status", "Decision", "Result",
           "Why it matters for the product", "Cross-team collaboration", "Main evidence", "Main link", "Cross-team link"]
WIDTHS = [5, 32, 15, 18, 11, 11, 13, 46, 40, 34, 44, 32, 10, 11]
TAB = re.compile(r"^(\d{4}) Highlights$")

FONT = "Arial"; PLUM = "4A1D4F"; GREY = "F5F5F5"; LINE = "D9D9D9"
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10); HDR_FILL = PatternFill("solid", fgColor=PLUM)
BODY = Font(name=FONT, size=10); BOLD = Font(name=FONT, size=10, bold=True)
LINK = Font(name=FONT, size=10, color="1F4E9E", underline="single")
H2 = Font(name=FONT, size=12, bold=True, color=PLUM)
WRAP = Alignment(wrap_text=True, vertical="top"); CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
THIN = Border(bottom=Side(style="thin", color=LINE))

GUIDE = [
    ("Year tabs", "Each year has its own tab. Rows are ranked by value to the product and to the users. The Status column shows what is shipped and what is still in progress."),
    ("Cross-team", "The \"Cross-team collaboration\" column shows the work I started or led with other teams, and the result. The \"Cross-team link\" column opens the issue or Slack thread."),
    ("Links", "\"Main link\" opens the main PR, issue or spec. The Slack links need access to the Terzo workspace. The GitHub links need access to the terzo-ai org."),
    ("Add a year", "Copy the newest year tab, rename it (for example \"2027 Highlights\"), clear the rows, and add a row for the year in the Years table. Change the formulas in the new row to use the new tab name."),
    ("Yellow cell", "Change \"Last updated\" each time you update the file."),
]


def iso(v):
    return v.date().isoformat() if isinstance(v, dt.datetime) else (v.isoformat() if isinstance(v, dt.date) else v)


def export(src, dst):
    wb = load_workbook(src)
    home = wb["Home"]
    data = {"title": home["B2"].value, "owner": home["B3"].value, "last_updated": iso(home["C5"].value),
            "notes": [], "years": []}
    in_notes = False
    for row in home.iter_rows(min_col=2, max_col=2):
        v = row[0].value
        if v == "Notes":
            in_notes = True
        elif in_notes and v:
            data["notes"].append(v)
    for ws in wb.worksheets:
        m = TAB.match(ws.title)
        if not m:
            continue
        rows = []
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=2).value is None:
                continue
            item = {f: iso(ws.cell(row=r, column=i).value) for i, f in enumerate(FIELDS, 1)}
            item["main_link"] = ws.cell(row=r, column=13).hyperlink.target if ws.cell(row=r, column=13).hyperlink else None
            item["cross_team_link"] = ws.cell(row=r, column=14).hyperlink.target if ws.cell(row=r, column=14).hyperlink else None
            rows.append(item)
        data["years"].append({"year": m.group(1), "rows": rows})
    data["years"].sort(key=lambda y: y["year"], reverse=True)
    json.dump(data, open(dst, "w"), indent=1, ensure_ascii=False)
    print(f"exported {sum(len(y['rows']) for y in data['years'])} rows in {len(data['years'])} year(s) to {dst}")


def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c); cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30


def year_tab(wb, year, rows):
    ws = wb.create_sheet(f"{year} Highlights"); ws.sheet_view.showGridLines = False
    header(ws, HEADERS, WIDTHS)
    rows = sorted(rows, key=lambda x: x["rank"])
    for i, item in enumerate(rows):
        r = i + 2
        for c, f in enumerate(FIELDS[:12], 1):
            v = item.get(f)
            if f in ("start", "latest") and v:
                v = dt.date.fromisoformat(v)
            cell = ws.cell(row=r, column=c, value=v); cell.font = BODY; cell.alignment = WRAP
        ws.cell(row=r, column=1).alignment = CENTER; ws.cell(row=r, column=1).font = Font(name=FONT, size=11, bold=True, color=PLUM)
        ws.cell(row=r, column=2).font = BOLD
        for c in (5, 6):
            ws.cell(row=r, column=c).number_format = "yyyy-mm-dd"; ws.cell(row=r, column=c).alignment = CENTER
        ws.cell(row=r, column=7).alignment = CENTER
        for c, f in ((13, "main_link"), (14, "cross_team_link")):
            cell = ws.cell(row=r, column=c); cell.alignment = CENTER
            if item.get(f):
                cell.value = "Open"; cell.hyperlink = item[f]; cell.font = LINK
        for c in range(1, len(HEADERS) + 1):
            ws.cell(row=r, column=c).border = THIN
            if i % 2 == 1:
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=GREY)
        ws.row_dimensions[r].height = 120
    n = max(len(rows) + 1, 2)
    ws.conditional_formatting.add(f"G2:G{n}", CellIsRule(operator="equal", formula=['"Shipped"'], fill=PatternFill("solid", fgColor="D9EAD3"), font=Font(name=FONT, size=10, bold=True, color="274E13")))
    ws.conditional_formatting.add(f"G2:G{n}", CellIsRule(operator="notEqual", formula=['"Shipped"'], fill=PatternFill("solid", fgColor="FFF2CC"), font=Font(name=FONT, size=10, bold=True, color="7F6000")))
    ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{n}"


def build(src, dst):
    data = json.load(open(src))
    years = sorted(data["years"], key=lambda y: y["year"], reverse=True)
    wb = Workbook()
    s = wb.active; s.title = "Home"; s.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFGH", [3, 22, 22, 14, 14, 14, 14, 14]):
        s.column_dimensions[col].width = w
    s["B2"] = data["title"]; s["B2"].font = Font(name=FONT, size=18, bold=True, color=PLUM)
    s["B3"] = data["owner"]; s["B3"].font = Font(name=FONT, size=10, color="595959")
    s["B5"] = "Last updated"; s["B5"].font = BOLD
    c = s["C5"]; c.value = dt.date.fromisoformat(data["last_updated"]); c.number_format = "mmmm d, yyyy"; c.font = BOLD
    c.fill = PatternFill("solid", fgColor="FFFF00"); c.alignment = Alignment(horizontal="left")
    c.comment = Comment("Change this date each time you update the file.", "Performance Highlights")
    s["B7"] = "Years"; s["B7"].font = H2
    for i, v in enumerate(["Year", "Tab", "Highlights", "Cross-team", "Shipped", "First date", "Latest date"], 2):
        cell = s.cell(row=8, column=i, value=v); cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    s.row_dimensions[8].height = 24
    r = 9
    for y in years:
        q = f"'{y['year']} Highlights'"
        vals = [y["year"], f"{y['year']} Highlights", f"=COUNTA({q}!B2:B100)", f"=COUNTA({q}!K2:K100)",
                f'=COUNTIF({q}!G2:G100,"Shipped")', f"=MIN({q}!E2:E100)", f"=MAX({q}!F2:F100)"]
        for i, v in enumerate(vals, 2):
            cell = s.cell(row=r, column=i, value=v); cell.font = BODY; cell.border = THIN; cell.alignment = Alignment(horizontal="left")
            if i >= 7:
                cell.number_format = "yyyy-mm-dd"
        s.cell(row=r, column=2).font = BOLD
        s.cell(row=r, column=3).hyperlink = f"#{q}!A1"; s.cell(row=r, column=3).font = LINK
        r += 1
    r += 1
    s.cell(row=r, column=2, value="How to read this file").font = H2
    for k, v in GUIDE:
        r += 1
        s.cell(row=r, column=2, value=k).font = BOLD; s.cell(row=r, column=2).alignment = WRAP
        cell = s.cell(row=r, column=3, value=v); cell.font = BODY; cell.alignment = WRAP
        s.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8); s.row_dimensions[r].height = 30
    r += 2
    s.cell(row=r, column=2, value="Notes").font = H2
    for n in data["notes"]:
        r += 1
        cell = s.cell(row=r, column=2, value=n); cell.font = Font(name=FONT, size=9, color="595959"); cell.alignment = WRAP
        s.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8); s.row_dimensions[r].height = 28
    for y in years:
        year_tab(wb, y["year"], y["rows"])
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = PLUM
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(dst)
    print(f"built {dst}: {', '.join(f'{y['year']} ({len(y['rows'])} rows)' for y in years)}, last updated {data['last_updated']}")


if __name__ == "__main__":
    if len(sys.argv) != 4 or sys.argv[1] not in ("export", "build"):
        sys.exit(__doc__)
    (export if sys.argv[1] == "export" else build)(sys.argv[2], sys.argv[3])
